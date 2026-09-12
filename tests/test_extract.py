"""M1 抽取管线的单测（requirements §7.5，D-15 / D-21 / D-25）。

夹具在 tmp_path 里现场生成：仓库不提交二进制测试文件（tests/fixtures 保持为空）。
生成 PDF 的文字一律用 ASCII —— 内置字体没有中文字形，中文会抽出乱码，
这是测试夹具的字体问题，不是抽取器的问题（真实作业书由 PyMuPDF 按原字形抽出）。
"""

import pytest

from src.intelligence.extract import (
    TABLE_MARKER,
    ExtractError,
    check_weight_sum,
    extract_text,
)
from src.models import RubricPoint


def _make_pdf(path, body, table):
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for line in body:
        page.insert_text(fitz.Point(72, y), line, fontsize=11)
        y += 16

    col_w, row_h, x0 = 130, 18, 72
    top = y + 12
    n_cols = max(len(row) for row in table)
    for r in range(len(table) + 1):
        yy = top + r * row_h
        page.draw_line(fitz.Point(x0, yy), fitz.Point(x0 + col_w * n_cols, yy))
    for c in range(n_cols + 1):
        xx = x0 + c * col_w
        page.draw_line(fitz.Point(xx, top), fitz.Point(xx, top + len(table) * row_h))
    for r, row in enumerate(table):
        for c, value in enumerate(row):
            page.insert_text(fitz.Point(x0 + c * col_w + 4, top + r * row_h + 13), value, fontsize=10)
    doc.save(str(path))
    doc.close()


def _make_docx(path, paragraphs, table=()):
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        t.style = "Table Grid"
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    doc.save(str(path))


def _point(pid, weight=None):
    return RubricPoint(id=pid, quote=f"{pid} 原文", observable="可核对", weight=weight)


# ---------- 分派与拒收 ----------


def test_text_passthrough(tmp_path):
    p = tmp_path / "brief.txt"
    p.write_text("作业书正文", encoding="utf-8")
    assert extract_text(p) == "作业书正文"


def test_image_is_rejected(tmp_path):
    p = tmp_path / "scan.png"
    p.write_bytes(b"not really an image")
    with pytest.raises(ExtractError) as exc:
        extract_text(p)
    assert "图片" in str(exc.value)


def test_unknown_suffix_is_rejected(tmp_path):
    p = tmp_path / "brief.pptx"
    p.write_bytes(b"x")
    with pytest.raises(ExtractError):
        extract_text(p)


def test_scanned_pdf_is_rejected(tmp_path):
    import fitz

    p = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(p))
    doc.close()
    with pytest.raises(ExtractError) as exc:
        extract_text(p)
    assert "文字层" in str(exc.value)


# ---------- PDF ----------


def test_pdf_body_and_table(tmp_path):
    p = tmp_path / "a.pdf"
    _make_pdf(p, ["Assignment brief"], [["Item", "Score"], ["Feature", "40"], ["Report", "60"]])
    text = extract_text(p)
    assert "Assignment brief" in text
    assert TABLE_MARKER in text
    assert "- Item：Feature｜Score：40" in text
    assert "- Item：Report｜Score：60" in text


def test_pdf_table_only_still_extractable(tmp_path):
    p = tmp_path / "t.pdf"
    _make_pdf(p, [], [["Item", "Score"], ["Feature", "100"]])
    text = extract_text(p)
    assert TABLE_MARKER in text
    assert "- Item：Feature｜Score：100" in text


# ---------- Word ----------


def test_docx_body_and_table(tmp_path):
    p = tmp_path / "a.docx"
    _make_docx(p, ["作业书正文"], [["评分项", "分值"], ["功能实现", "40"], ["报告", "60"]])
    text = extract_text(p)
    assert "作业书正文" in text
    assert TABLE_MARKER in text
    assert "- 评分项：功能实现｜分值：40" in text


def test_docx_empty_cells_are_skipped(tmp_path):
    p = tmp_path / "b.docx"
    _make_docx(p, [], [["评分项", "分值"], ["功能实现", ""]])
    text = extract_text(p)
    assert "- 评分项：功能实现" in text
    assert "分值：" not in text


def test_empty_docx_is_rejected(tmp_path):
    p = tmp_path / "e.docx"
    _make_docx(p, [])
    with pytest.raises(ExtractError):
        extract_text(p)


# ---------- 权重加总软校验（§7.5 / D-25）----------


def test_weight_sum_in_range_is_silent():
    assert check_weight_sum([_point("R1", 40), _point("R2", 60)]) is None
    assert check_weight_sum([_point("R1", 95)]) is None


def test_weight_sum_missing_is_skipped_not_warned():
    assert check_weight_sum([_point("R1", 60), _point("R2")]) is None
    assert check_weight_sum([_point("R1")]) is None


def test_weight_sum_twenty_scale_is_skipped():
    # 20 分制：能归一化到百分制，但不是百分制量纲 → 整体跳过，不误报
    assert check_weight_sum([_point("R1", 10), _point("R2", 10)]) is None


def test_weight_sum_lost_rows_warns():
    warning = check_weight_sum([_point("R1", 40), _point("R2", 45)])
    assert warning is not None
    assert "丢行" in warning
    assert "85" in warning


def test_weight_sum_over_hundred_warns():
    assert check_weight_sum([_point("R1", 70), _point("R2", 50)]) is not None


def test_weight_sum_of_empty_rubric_is_silent():
    assert check_weight_sum([]) is None

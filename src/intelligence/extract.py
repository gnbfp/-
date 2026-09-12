"""M1 前置步骤：文件 → 纯文本（**不涉及 LLM**）。

依据：requirements.md §7.5、docs/ARCHITECTURE.md §6.3、D-15 / D-21 / D-25。

产物形态就是一个 ``{ text: string }``：正文 + 表格区块。表格用
``find_tables()`` 单独抽，与正文用 ``=== 表格区域 ===`` 分隔；每行
``- 列名：值｜列名：值``（行间换行，避免跨行粘连 —— 待定义-33）。

职责边界：
  * 本模块**不调 LLM**、不落盘、不判定评分点。
  * ``check_weight_sum()`` 的入参是 **M1 解析之后**的 ``RubricPoint`` 列表
    （``weight`` 是 LLM 抽出来的）。它按 §6.3 的管线顺序放在这里，属于软校验：
    命中只返回一句警告，**绝不拒收**。
"""

from __future__ import annotations

from pathlib import Path
from typing import Sequence

from src.models import RubricPoint

__all__ = [
    "ExtractError",
    "TABLE_MARKER",
    "CELL_SEP",
    "VALUE_SEP",
    "ROW_PREFIX",
    "extract_text",
    "check_weight_sum",
]

TABLE_MARKER = "=== 表格区域 ==="
CELL_SEP = "｜"           # 列内分隔（待定义-33）
VALUE_SEP = "："          # 列名：值
ROW_PREFIX = "- "         # 每行前缀（待定义-33）

_TEXT_SUFFIXES = {"", ".txt", ".md", ".markdown", ".text"}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".heic"}

WEIGHT_TARGET = 100.0     # 百分制满分
WEIGHT_TOLERANCE = 10.0   # 加总落在 [90, 110] 视为正常（§7.5）
# "能归一化到百分制"的落地口径：总分太小的量纲（如 20 分制）直接跳过，不误报。
# 之所以需要这道门：丢了几行的百分制评分表（如加总 85）与 20 分制在数值上无法
# 仅凭加总区分；低于本门限一律视为另一种量纲。此取值为工程默认，待真实作业书校准。
PERCENT_SCALE_MIN = 50.0


class ExtractError(RuntimeError):
    """文件抽不出文本 —— 扫描版 PDF / 图片 / 不认识的格式。消息直接可以发给用户。"""


def extract_text(path: Path | str) -> str:
    """按后缀分派：PDF / Word / 纯文本。图片与扫描版直接拒收（§7.5）。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return _pdf_to_text(p)
    if suffix == ".docx":
        return _docx_to_text(p)
    if suffix in _TEXT_SUFFIXES:
        return p.read_text(encoding="utf-8")
    if suffix in _IMAGE_SUFFIXES:
        raise ExtractError(f"图片不能抽取文字（OCR 已砍，P2）：{p.name}。请发文字版作业书。")
    raise ExtractError(f"不认识的文件类型 {suffix or '(无后缀)'}：{p.name}。支持 PDF / DOCX / TXT。")


# ---------- PDF ----------


def _pdf_to_text(path: Path) -> str:
    import fitz  # PyMuPDF：懒导入，纯文本链路不因它缺失而 import 失败

    body_chunks: list[str] = []
    table_lines: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            body_chunks.append(page.get_text("text"))
            for table in page.find_tables().tables:
                names, data = _split_header(table)
                table_lines.extend(_rows_to_lines(data, names))

    body = "\n".join(body_chunks).strip()
    if not body and not table_lines:
        raise ExtractError(
            f"PDF 没有文字层（扫描版或图片版）：{path.name}。请发文字版作业书。"
        )
    return _assemble(body, table_lines)


def _split_header(table) -> tuple[list[str] | None, list[list[str]]]:
    """拆出列名与数据行。PyMuPDF 的 ``header.external`` 决定表头是否已在 extract() 里。"""
    rows = [[_cell(c) for c in row] for row in table.extract()]
    rows = [row for row in rows if any(row)]
    if not rows:
        return None, []
    try:
        names = [_cell(n) for n in table.header.names]
        external = bool(table.header.external)
    except Exception:                      # 低版本 / 未识别表头：退回"首行即表头"
        names, external = None, False
    if not names or not any(names):
        return rows[0], rows[1:]
    return names, rows if external else rows[1:]


# ---------- Word ----------


def _docx_to_text(path: Path) -> str:
    from docx import Document  # python-docx

    doc = Document(path)
    body = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    table_lines: list[str] = []
    for table in doc.tables:
        rows = [[_cell(c.text) for c in row.cells] for row in table.rows]
        rows = [row for row in rows if any(row)]
        if not rows:
            continue
        table_lines.extend(_rows_to_lines(rows[1:], rows[0]))

    if not body and not table_lines:
        raise ExtractError(f"Word 文档没有可抽取的文字：{path.name}。")
    return _assemble(body, table_lines)


# ---------- 共享格式 ----------


def _cell(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())      # 压掉单元格里的换行与连续空白


def _rows_to_lines(rows: Sequence[Sequence[str]], names: Sequence[str] | None) -> list[str]:
    """数据行 → ``- 列名：值｜列名：值``；空单元格跳过。"""
    lines: list[str] = []
    for row in rows:
        pairs = []
        for index, value in enumerate(row):
            if not value:
                continue
            key = names[index] if names and index < len(names) and names[index] else f"列{index + 1}"
            pairs.append(f"{key}{VALUE_SEP}{value}")
        if pairs:
            lines.append(ROW_PREFIX + CELL_SEP.join(pairs))
    return lines


def _assemble(body: str, table_lines: Sequence[str]) -> str:
    parts = [body] if body else []
    if table_lines:
        parts.append(TABLE_MARKER)
        parts.extend(table_lines)
    return "\n".join(parts)


# ---------- 权重加总软校验（§7.5，原待定义-32 / D-25）----------


def check_weight_sum(points: Sequence[RubricPoint]) -> str | None:
    """权重加总软校验。返回警告文案；``None`` = 跳过或正常。

    前置条件（缺一即整体跳过，不误报）：
      * 每一个评分点都有 ``weight``；
      * 加总为正，且量纲是百分制（``>= PERCENT_SCALE_MIN``，20 分制之类跳过）。

    **软警告不拒收**：调用方（M1）把返回值拼进给用户的回复即可。
    """
    if not points:
        return None
    weights = [p.weight for p in points]
    if any(weight is None for weight in weights):
        return None
    total = float(sum(weights))              # type: ignore[arg-type]
    if total <= 0 or total < PERCENT_SCALE_MIN:
        return None
    low, high = WEIGHT_TARGET - WEIGHT_TOLERANCE, WEIGHT_TARGET + WEIGHT_TOLERANCE
    if low <= total <= high:
        return None
    return (
        f"评分标准解析疑似丢行，请对照原文核对"
        f"（评分点权重加总 = {total:g}，应接近 {WEIGHT_TARGET:g}）"
    )
